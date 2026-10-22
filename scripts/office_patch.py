#!/usr/bin/env python3
"""office_patch.py — 克劳德拉做申报材料时攒下来的两个小工具。

子命令
------
swap-docx-image
    就地替换 docx 里某个内嵌图片（word/media/imageN.ext），并可选把新图裁成
    与文档里显示框相同的宽高比，避免被拉伸变形。

fit-pptx-titles
    把所有“页眉标题”文本框统一加宽（默认 14.2 in），关掉自动换行与自动缩放，
    杜绝标题被挤成两行。

用法示例
--------
    python3 office_patch.py swap-docx-image \
        --docx 省自科申报书.docx \
        --media word/media/image13.jpeg \
        --image 新截图.jpg \
        --match-aspect

    python3 office_patch.py fit-pptx-titles --pptx deck.pptx --width-in 14.2

两条命令都在原文件上工作，替换前请自行备份；加 --dry-run 只看结果不动文件。
"""

import argparse
import re
import shutil
import sys
import zipfile
from io import BytesIO

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


# ---------------------------------------------------------------- docx 图片替换

def _rid_for_target(rels_xml: bytes, target_suffix: str) -> str | None:
    """在 word/_rels/document.xml.rels 里找到指向该 media 的 rId。"""
    from lxml import etree

    root = etree.fromstring(rels_xml)
    for rel in root.findall(f"{{{REL_NS}}}Relationship"):
        target = rel.get("Target", "")
        if target.replace("\\", "/").endswith(target_suffix.replace("word/", "")):
            return rel.get("Id")
    return None


def _extent_for_rid(document_xml: bytes, rid: str) -> tuple[int, int] | None:
    """找到引用该 rId 的 drawing，返回它的 wp:extent(cx, cy)。"""
    from lxml import etree

    root = etree.fromstring(document_xml)
    for container in root.iter(f"{{{WP_NS}}}inline", f"{{{WP_NS}}}anchor"):
        for blip in container.iter(f"{{{A_NS}}}blip"):
            if blip.get(f"{{{R_NS}}}embed") == rid:
                extent = container.find(f"{{{WP_NS}}}extent")
                if extent is not None:
                    return int(extent.get("cx")), int(extent.get("cy"))
    return None


def _fit_aspect(image_path: str, cx: int, cy: int, target_format: str):
    """把图裁成 cx/cy 的比例：从顶部裁掉多余部分（截图顶部多为留白）。"""
    from PIL import Image

    im = Image.open(image_path).convert("RGB")
    w, h = im.size
    want_h = round(w * cy / cx)
    if want_h <= h:
        im = im.crop((0, h - want_h, w, h))
        note = f"从顶部裁掉 {h - want_h}px"
    else:
        # 高度不够，改从左右裁
        want_w = round(h * cx / cy)
        left = (w - want_w) // 2
        im = im.crop((left, 0, left + want_w, h))
        note = f"从两侧各裁 {left}px"
    buf = BytesIO()
    fmt = "JPEG" if target_format.lower().endswith(("jpg", "jpeg")) else "PNG"
    if fmt == "JPEG":
        im.save(buf, fmt, quality=95, optimize=True)
    else:
        im.save(buf, fmt)
    return buf.getvalue(), im.size, note


def cmd_swap_docx_image(args) -> int:
    with zipfile.ZipFile(args.docx) as z:
        names = z.namelist()
        if args.media not in names:
            print(f"docx 里没有 {args.media}，现有 media：")
            for n in names:
                if n.startswith("word/media/"):
                    print("   ", n)
            return 1
        document = z.read("word/document.xml")
        rels = z.read("word/_rels/document.xml.rels")

        rid = _rid_for_target(rels, args.media)
        if rid is None:
            print("没在 rels 里找到这条 media 的关系，可能路径不对")
            return 1
        extent = _extent_for_rid(document, rid)
        if extent is None:
            print(f"图 {args.media}（{rid}）在正文里没有被引用，跳过宽高比处理")
        data: bytes
        if args.match_aspect and extent:
            cx, cy = extent
            data, size, note = _fit_aspect(args.image, cx, cy, args.media)
            print(f"显示框 {cx}x{cy} EMU（比例 {cx / cy:.4f}），新图 {size[0]}x{size[1]}，{note}")
        else:
            data = open(args.image, "rb").read()
            print(f"直接替换 {args.media}，{len(data)} 字节")

        if args.dry_run:
            print("--dry-run：不写入")
            return 0

        tmp = args.docx + ".tmp"
        with zipfile.ZipFile(args.docx) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                payload = data if item.filename == args.media else zin.read(item.filename)
                zout.writestr(item, payload)
        shutil.move(tmp, args.docx)
        print(f"已写回 {args.docx}")
    return 0


# ---------------------------------------------------------------- pptx 标题宽度

def cmd_fit_pptx_titles(args) -> int:
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches

    prs = Presentation(args.pptx)
    changed = []
    for idx, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame or shape.left is None or shape.top is None:
                continue
            if abs(shape.left / 914400 - args.left_in) > 0.08:
                continue
            if shape.top / 914400 > args.top_max_in:
                continue
            if not shape.text_frame.text.strip():
                continue
            before = shape.width / 914400 if shape.width else 0
            shape.width = Inches(args.width_in)
            shape.text_frame.word_wrap = False
            shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
            changed.append((idx, shape.text_frame.text.strip()[:18], round(before, 2)))
    for idx, text, before in changed:
        print(f"  第{idx:>2}页 标题 {text!r} 宽 {before} → {args.width_in} in")
    print(f"共处理 {len(changed)} 个标题框")
    if args.dry_run:
        print("--dry-run：不写入")
        return 0
    prs.save(args.pptx)
    print(f"已写回 {args.pptx}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p1 = sub.add_parser("swap-docx-image", help="替换 docx 内嵌图片")
    p1.add_argument("--docx", required=True)
    p1.add_argument("--media", required=True, help="如 word/media/image13.jpeg")
    p1.add_argument("--image", required=True, help="新图路径")
    p1.add_argument("--match-aspect", action="store_true", help="裁成文档里显示框的宽高比")
    p1.add_argument("--dry-run", action="store_true")
    p1.set_defaults(func=cmd_swap_docx_image)

    p2 = sub.add_parser("fit-pptx-titles", help="统一 PPT 页眉标题框宽度")
    p2.add_argument("--pptx", required=True)
    p2.add_argument("--left-in", type=float, default=2.0, help="标题左边界（英寸），默认 2.0")
    p2.add_argument("--top-max-in", type=float, default=1.2, help="顶部阈值（英寸），默认 1.2")
    p2.add_argument("--width-in", type=float, default=14.2, help="目标宽度（英寸），默认 14.2")
    p2.add_argument("--dry-run", action="store_true")
    p2.set_defaults(func=cmd_fit_pptx_titles)

    args = ap.parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
